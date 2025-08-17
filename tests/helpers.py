import hashlib
import io
import re
from collections.abc import Sequence
from datetime import UTC, datetime

import docx
import numpy as np
from aiogram.client.session.base import BaseSession
from aiogram.methods import (
    EditMessageText,
    GetChatMember,
    GetFile,
    GetMe,
    SendMessage,
    TelegramMethod,
)
from aiogram.types import Chat, ChatMemberMember, ChatMemberOwner, File, Message, User


class FakeEmbedder:
    """Hashed bag of words. Texts sharing words get similar vectors, which is enough
    to make retrieval tests meaningful without downloading a model."""

    dimension = 256

    def __init__(self) -> None:
        self.passage_calls = 0

    def _vector(self, text: str) -> list[float]:
        vec = np.zeros(self.dimension, dtype=np.float32)
        for word in re.findall(r"\w+", text.lower()):
            bucket = int(hashlib.md5(word.encode()).hexdigest(), 16) % self.dimension
            vec[bucket] += 1.0
        norm = np.linalg.norm(vec)
        if norm:
            vec /= norm
        return vec.tolist()

    async def embed_passages(self, texts: Sequence[str]) -> list[list[float]]:
        self.passage_calls += 1
        return [self._vector(t) for t in texts]

    async def embed_query(self, text: str) -> list[float]:
        return self._vector(text)


class FakeLLM:
    def __init__(self, reply: str = "Ответ [1].", error: Exception | None = None) -> None:
        self.reply = reply
        self.error = error
        self.calls: list[list[dict[str, str]]] = []

    async def complete(self, messages: list[dict[str, str]]) -> str:
        self.calls.append(messages)
        if self.error is not None:
            raise self.error
        return self.reply


def make_pdf(pages: list[list[str]]) -> bytes:
    """Build a minimal text PDF. Each page is a list of lines (Latin-1 only)."""

    def esc(line: str) -> str:
        return line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")

    objects: list[bytes] = []
    page_refs = []
    font_id = 3
    next_id = 4
    page_objects: list[tuple[int, bytes]] = []
    for lines in pages:
        page_id, content_id = next_id, next_id + 1
        next_id += 2
        page_refs.append(f"{page_id} 0 R")
        ops = ["BT", "/F1 12 Tf", "14 TL", "72 720 Td"]
        for i, line in enumerate(lines):
            ops.append(f"({esc(line)}) Tj" if i == 0 else f"T* ({esc(line)}) Tj")
        ops.append("ET")
        stream = "\n".join(ops).encode("latin-1") if lines else b""
        page_objects.append(
            (
                page_id,
                f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                f"/Resources << /Font << /F1 {font_id} 0 R >> >> "
                f"/Contents {content_id} 0 R >>".encode(),
            )
        )
        page_objects.append(
            (
                content_id,
                b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
            )
        )

    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objects.append(f"<< /Type /Pages /Kids [{' '.join(page_refs)}] /Count {len(pages)} >>".encode())
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    objects.extend(body for _, body in sorted(page_objects))

    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % number + body + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1))
    for offset in offsets:
        out.write(b"%010d 00000 n \n" % offset)
    out.write(
        b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref)
    )
    return out.getvalue()


def make_docx(blocks: list[tuple[str, str]]) -> bytes:
    """blocks: ("h1" | "p", text) or ("table", "a|b;c|d")."""
    document = docx.Document()
    for kind, text in blocks:
        if kind == "h1":
            document.add_heading(text, level=1)
        elif kind == "table":
            rows = [row.split("|") for row in text.split(";")]
            table = document.add_table(rows=len(rows), cols=len(rows[0]))
            for r, row in enumerate(rows):
                for c, value in enumerate(row):
                    table.cell(r, c).text = value
        else:
            document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class MockedSession(BaseSession):
    """Records every Bot API call and answers with plausible objects instead of hitting Telegram."""

    def __init__(self, files: dict[str, bytes] | None = None) -> None:
        super().__init__()
        self.requests: list[TelegramMethod] = []
        self.files = files if files is not None else {}
        self.admins: set[int] = set()
        self._message_id = 1000

    async def make_request(self, bot, method, timeout=None):
        self.requests.append(method)
        if isinstance(method, SendMessage | EditMessageText):
            self._message_id += 1
            return Message(
                message_id=getattr(method, "message_id", None) or self._message_id,
                date=datetime.now(UTC),
                chat=Chat(id=method.chat_id, type="private"),
                text=method.text,
            ).as_(bot)
        if isinstance(method, GetFile):
            return File(file_id=method.file_id, file_unique_id="u", file_path=method.file_id)
        if isinstance(method, GetChatMember):
            member = User(id=method.user_id, is_bot=False, first_name="U")
            if method.user_id in self.admins:
                return ChatMemberOwner(user=member, is_anonymous=False)
            return ChatMemberMember(user=member)
        if isinstance(method, GetMe):
            return User(id=42, is_bot=True, first_name="AskDocs", username="askdocs_test_bot")
        return True

    async def stream_content(
        self, url, headers=None, timeout=30, chunk_size=65536, raise_for_status=True
    ):
        yield self.files[url.rsplit("/", 1)[-1]]

    async def close(self) -> None:
        pass

    def texts(self) -> list[str]:
        return [r.text for r in self.requests if isinstance(r, SendMessage | EditMessageText)]
